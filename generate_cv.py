import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def create_element(name):
    return OxmlElement(name)

def set_cell_margins(cell, top=70, bottom=70, left=90, right=90):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_background(cell, color_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
    cell._element.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_table_borders(table, color="CBD5E1", sz="4"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders_xml = f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
            <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
        '''
        tblPr[0].append(parse_xml(borders_xml))

def add_section_header(doc, title_text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(11)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    
    run = p.add_run(title_text.upper())
    run.font.name = 'Calibri'
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = RGBColor(15, 76, 129)  # Deep Academic Navy
    
    # Elegant bottom border line
    pPr = p._p.get_or_add_pPr()
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="8" w:space="2" w:color="0F4C81"/></w:pBdr>')
    pPr.append(pBdr)
    return p

def set_classic_grid_borders(table, color="94A3B8", sz="4"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders_xml = f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
        </w:tblBorders>
        '''
        tblPr[0].append(parse_xml(borders_xml))

def add_classic_section_header(doc, title_text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    
    # Classic banner with shaded background and navy border
    pPr = p._p.get_or_add_pPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="E2E8F0"/>')
    pPr.append(shd)
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="2" w:color="0F4C81"/></w:pBdr>')
    pPr.append(pBdr)
    
    run = p.add_run(f"  {title_text.upper()}")
    run.font.name = 'Calibri'
    run.font.size = Pt(9.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(15, 76, 129)
    return p

def build_best_teacher_cv():
    doc = Document()
    
    # Balanced A4 Margins (Consistent aesthetic from our best version)
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    
    # Global Default Font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(9.5)
    font.color.rgb = RGBColor(30, 41, 59) # Slate 800
    
    # =========================================================================
    # HEADER TABLE (Name, Designation, Contact & Photo)
    # =========================================================================
    header_table = doc.add_table(rows=1, cols=2)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False
    
    col_widths = [Inches(5.4), Inches(1.55)]
    for row in header_table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width
            
    info_cell = header_table.rows[0].cells[0]
    photo_cell = header_table.rows[0].cells[1]
    
    set_cell_margins(info_cell, top=0, bottom=30, left=0, right=70)
    set_cell_margins(photo_cell, top=0, bottom=30, left=30, right=0)
    photo_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    
    # Name
    p_name = info_cell.paragraphs[0]
    p_name.paragraph_format.space_after = Pt(2)
    p_name.paragraph_format.space_before = Pt(0)
    r_name = p_name.add_run("NURUL AREFIN NABIL")
    r_name.font.name = 'Calibri'
    r_name.font.size = Pt(20)
    r_name.font.bold = True
    r_name.font.color.rgb = RGBColor(15, 76, 129)
    
    # Teacher Designation
    p_title = info_cell.add_paragraph()
    p_title.paragraph_format.space_after = Pt(5)
    r_title = p_title.add_run("ICT & Computer Science Teacher | B.Sc. in CSE")
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(11)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(16, 115, 95) # Deep Academic Teal
    
    # Contact Details
    contacts = [
        ("📍 Address: ", "Barmi, Sreepur, Gazipur-1743, Bangladesh"),
        ("📞 Phone: ", "+880 1881-196146"),
        ("✉ Email: ", "nurularefinnabil@gmail.com"),
        ("🌐 Portfolio: ", "www.areefin.me   |   LinkedIn: in/n-arefin-nabil   |   GitHub: arefin-nabil")
    ]
    
    for lbl, val in contacts:
        p_c = info_cell.add_paragraph()
        p_c.paragraph_format.space_after = Pt(1)
        p_c.paragraph_format.line_spacing = 1.05
        r_lbl = p_c.add_run(lbl)
        r_lbl.font.bold = True
        r_lbl.font.size = Pt(8.8)
        r_lbl.font.color.rgb = RGBColor(71, 85, 105)
        r_val = p_c.add_run(val)
        r_val.font.size = Pt(8.8)
        r_val.font.color.rgb = RGBColor(30, 41, 59)
        
    # Photo insertion
    img_path = "profile.jpg"
    p_photo = photo_cell.paragraphs[0]
    p_photo.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_photo.paragraph_format.space_after = Pt(0)
    if os.path.exists(img_path):
        p_photo.add_run().add_picture(img_path, width=Inches(1.35))
        
    # =========================================================================
    # 1. CAREER OBJECTIVE
    # =========================================================================
    add_section_header(doc, "1. Career Objective")
    p_obj = doc.add_paragraph()
    p_obj.paragraph_format.space_after = Pt(4)
    p_obj.paragraph_format.line_spacing = 1.15
    p_obj.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r_obj = p_obj.add_run(
        "To work as an ICT Teacher where I can use my academic background, teaching experience, and practical computer skills "
        "to deliver clear, engaging lessons and help students develop strong ICT knowledge, programming skills, "
        "and problem-solving abilities through effective classroom and lab-based teaching."
    )
    r_obj.font.size = Pt(9.5)
    
    # =========================================================================
    # 2. TEACHING EXPERIENCE
    # =========================================================================
    add_section_header(doc, "2. Teaching Experience")
    
    # Role 1: Instructor / Tutor
    p_exp1 = doc.add_paragraph()
    p_exp1.paragraph_format.space_before = Pt(3)
    p_exp1.paragraph_format.space_after = Pt(1)
    r1 = p_exp1.add_run("ICT & Programming Instructor / Academic Tutor")
    r1.font.bold = True
    r1.font.size = Pt(9.8)
    r1.font.color.rgb = RGBColor(15, 76, 129)
    p_exp1.add_run("   |   Tongi, Uttara & Sreepur").font.size = Pt(9)
    r1_date = p_exp1.add_run(" | 2024 – Present (1+ Year)")
    r1_date.font.italic = True
    r1_date.font.size = Pt(8.8)
    r1_date.font.color.rgb = RGBColor(100, 116, 139)
    
    exp1_bullets = [
        "Delivered structured lessons on **HSC ICT (National Curriculum)**, specializing in **C Programming (Ch 5)**, **HTML/Web Design (Ch 4)**, **Database/SQL (Ch 6)**, and Number Systems.",
        "Conducted hands-on practical lab sessions, guiding students to write, debug, and execute code using Code::Blocks and text editors.",
        "Used previous years' board questions, model tests, and structured worksheets to support exam preparation and improve problem-solving skills.",
        "Integrated digital teaching tools and modern AI-assisted study methods to simplify abstract algorithmic concepts."
    ]
    for b in exp1_bullets:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(1)
        bp.paragraph_format.line_spacing = 1.08
        parts = b.split("**")
        for i, part in enumerate(parts):
            r = bp.add_run(part)
            r.font.size = Pt(8.8)
            if i % 2 == 1:
                r.font.bold = True

    # Role 2: Assistant Teacher
    p_exp2 = doc.add_paragraph()
    p_exp2.paragraph_format.space_before = Pt(4)
    p_exp2.paragraph_format.space_after = Pt(1)
    r2 = p_exp2.add_run("Assistant Teacher (ICT & Science)")
    r2.font.bold = True
    r2.font.size = Pt(9.8)
    r2.font.color.rgb = RGBColor(15, 76, 129)
    p_exp2.add_run("   |   Barmi Al-Madina School & Plus Coaching").font.size = Pt(9)
    r2_date = p_exp2.add_run(" | 1+ Year")
    r2_date.font.italic = True
    r2_date.font.size = Pt(8.8)
    r2_date.font.color.rgb = RGBColor(100, 116, 139)
    
    exp2_bullets = [
        "**Classroom Management:** Maintained a supportive and disciplined classroom environment and addressed individual student learning needs.",
        "**Lesson Planning:** Prepared structured lesson plans, class notes, worksheets, and practice materials aligned with curriculum standards.",
        "**Assessment & Evaluation:** Conducted periodic class tests, practical lab quizzes, and model exams; provided constructive feedback to boost student confidence.",
        "**Lab Supervision:** Supervised computer lab sessions, trained students in touch-typing, MS Office applications, and safe internet practices."
    ]
    for b in exp2_bullets:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(1)
        bp.paragraph_format.line_spacing = 1.08
        parts = b.split("**")
        for i, part in enumerate(parts):
            r = bp.add_run(part)
            r.font.size = Pt(8.8)
            if i % 2 == 1:
                r.font.bold = True

    # =========================================================================
    # 3. EDUCATIONAL QUALIFICATIONS
    # =========================================================================
    add_section_header(doc, "3. Educational Qualifications")
    
    edu_table = doc.add_table(rows=4, cols=5)
    edu_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    edu_table.autofit = False
    set_table_borders(edu_table, color="CBD5E1", sz="4")
    
    edu_col_widths = [Inches(1.75), Inches(2.3), Inches(1.3), Inches(0.8), Inches(0.82)]
    for row in edu_table.rows:
        for idx, width in enumerate(edu_col_widths):
            row.cells[idx].width = width
            
    headers = ["Exam / Degree", "Institution / Board", "Major / Group", "Year", "Result"]
    for idx, text in enumerate(headers):
        cell = edu_table.rows[0].cells[idx]
        set_cell_background(cell, "F1F5F9")
        set_cell_margins(cell, top=65, bottom=65, left=65, right=65)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        if idx in [3, 4]:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.font.bold = True
        r.font.size = Pt(8.8)
        r.font.color.rgb = RGBColor(15, 76, 129)
        
    edu_data = [
        ("B.Sc. in Computer Science & Engineering (CSE)", "Uttara University, Dhaka", "Computer Science", "2026", "CGPA: 3.40\n/ 4.00"),
        ("Alim (Equivalent to HSC)", "Tamirul Millat Kamil Madrasha\n(Madrasah Education Board)", "Science Group", "2021", "GPA: 5.00\n/ 5.00"),
        ("Dakhil (Equivalent to SSC)", "Patka Dakhil Madrasha\n(Madrasah Education Board)", "Science Group", "2019", "GPA: 5.00\n(Upazila 1st)")
    ]
    
    for row_idx, data in enumerate(edu_data, start=1):
        for col_idx, text in enumerate(data):
            cell = edu_table.rows[row_idx].cells[col_idx]
            set_cell_margins(cell, top=55, bottom=55, left=65, right=65)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if col_idx in [3, 4]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(text)
            r.font.size = Pt(8.6)
            if col_idx == 0:
                r.font.bold = True
            if col_idx == 4 and "5.00" in text:
                r.font.bold = True

    # =========================================================================
    # 4. TEACHING COMPETENCIES & CURRICULUM EXPERTISE
    # =========================================================================
    add_section_header(doc, "4. Teaching Competencies & Curriculum Expertise")
    
    comp_items = [
        ("Curriculum Knowledge:", "Comprehensive grasp of National Curriculum HSC ICT (Chapters 1–6) and Secondary ICT."),
        ("Programming & Web:", "Strong command in teaching C Programming fundamentals, algorithm design, HTML & CSS web basics."),
        ("Database & Systems:", "Relational Database Management Concepts, SQL queries, Number Systems, and Digital Logic."),
        ("Laboratory & Practical:", "Hands-on computer lab conducting, typing training, software installation, and hardware troubleshooting."),
        ("Pedagogy & Classroom:", "Structured lesson planning, customized worksheets, continuous student assessment, and board exam model tests."),
        ("Digital Instruction:", "Smart classroom technologies, multimedia slide design, and technology-assisted teaching methodologies.")
    ]
    for c_lbl, c_val in comp_items:
        p_c = doc.add_paragraph()
        p_c.paragraph_format.space_after = Pt(1.5)
        p_c.paragraph_format.line_spacing = 1.05
        r_l = p_c.add_run(f"•  {c_lbl} ")
        r_l.font.bold = True
        r_l.font.size = Pt(8.8)
        r_l.font.color.rgb = RGBColor(15, 76, 129)
        r_v = p_c.add_run(c_val)
        r_v.font.size = Pt(8.6)

    # =========================================================================
    # 5. ICT & COMPUTER SKILLS
    # =========================================================================
    add_section_header(doc, "5. ICT & Computer Skills")
    
    skills_data = [
        ("ICT & Office Applications: ", "MS Word, MS Excel, MS PowerPoint, MS Access, Internet & Email, Google Workspace"),
        ("Programming Languages: ", "C, C++, Python, Java, Basic Web Development (HTML5, CSS3)"),
        ("Database & Data: ", "MySQL, SQLite, Relational Database Concepts & Basic SQL"),
        ("Computing & Hardware: ", "Operating Systems (Windows/Linux), Hardware Troubleshooting, Basic LAN Networking"),
        ("Teaching Technology: ", "Smart Classroom Systems, Multimedia Projectors, Interactive Online Learning Tools"),
        ("Additional Technical Stack: ", "Flutter/Dart, REST API, VS Code, Code::Blocks, Git/GitHub")
    ]
    for k, v in skills_data:
        p_sk = doc.add_paragraph()
        p_sk.paragraph_format.space_after = Pt(1.5)
        p_sk.paragraph_format.line_spacing = 1.05
        r_k = p_sk.add_run(f"•  {k}")
        r_k.font.bold = True
        r_k.font.size = Pt(8.8)
        r_v = p_sk.add_run(v)
        r_v.font.size = Pt(8.6)

    # =========================================================================
    # 6. SELECTED PRACTICAL & RESEARCH PROJECTS
    # =========================================================================
    add_section_header(doc, "6. Selected Practical & Research Projects")
    
    # Project 1: Snake App
    p_pr1 = doc.add_paragraph()
    p_pr1.paragraph_format.space_before = Pt(2.5)
    p_pr1.paragraph_format.space_after = Pt(1)
    r_p1 = p_pr1.add_run("1. Snake & Wildlife BD – Android Application")
    r_p1.font.bold = True
    r_p1.font.size = Pt(9.2)
    r_p1.font.color.rgb = RGBColor(15, 76, 129)
    r_p1_badge = p_pr1.add_run("   |   Published on Google Play Store")
    r_p1_badge.font.bold = True
    r_p1_badge.font.size = Pt(8.6)
    r_p1_badge.font.color.rgb = RGBColor(16, 115, 95)
    
    p_pr1_desc = doc.add_paragraph()
    p_pr1_desc.paragraph_format.space_after = Pt(2)
    p_pr1_desc.paragraph_format.line_spacing = 1.08
    p_pr1_desc.paragraph_format.left_indent = Inches(0.18)
    r_p1_d = p_pr1_desc.add_run(
        "• Stack & Scope: Flutter, Dart, REST API, Cloud Database.\n"
        "• Overview: Developed and published an educational and wildlife awareness mobile application, integrating digital learning materials, emergency helpline connectivity, and nationwide rescuer directory for instant public safety."
    )
    r_p1_d.font.size = Pt(8.6)
    
    # Project 2: Fake News Detection
    p_pr2 = doc.add_paragraph()
    p_pr2.paragraph_format.space_before = Pt(3)
    p_pr2.paragraph_format.space_after = Pt(1)
    r_p2 = p_pr2.add_run("2. Bangla Fake News Detection – Machine Learning Research Project")
    r_p2.font.bold = True
    r_p2.font.size = Pt(9.2)
    r_p2.font.color.rgb = RGBColor(15, 76, 129)
    r_p2_badge = p_pr2.add_run("   |   Undergraduate Research")
    r_p2_badge.font.bold = True
    r_p2_badge.font.size = Pt(8.6)
    r_p2_badge.font.color.rgb = RGBColor(16, 115, 95)
    
    p_pr2_desc = doc.add_paragraph()
    p_pr2_desc.paragraph_format.space_after = Pt(2)
    p_pr2_desc.paragraph_format.line_spacing = 1.08
    p_pr2_desc.paragraph_format.left_indent = Inches(0.18)
    r_p2_d = p_pr2_desc.add_run(
        "• Stack & Scope: Python, Scikit-Learn, NLTK, TF-IDF Vectorization, Machine Learning.\n"
        "• Overview: Developing a Natural Language Processing (NLP) classification model to identify and flag fabricated online news articles by analyzing textual veracity, linguistic features, and lexical patterns."
    )
    r_p2_d.font.size = Pt(8.6)

    # =========================================================================
    # 7. ACADEMIC HONORS & TRAINING
    # =========================================================================
    add_section_header(doc, "7. Academic Honors & Training")
    
    honors_trainings = [
        ("Upazila 1st Rank (Dakhil Exam, 2019): ", "Secured 1st Position in Sreepur Upazila under Bangladesh Madrasah Board with GPA 5.00/5.00."),
        ("General Board Merit Scholarship: ", "Awarded merit-based scholarship by the Ministry of Education for academic excellence."),
        ("Competitive Programming & Problem Solving: ", "Completed structured algorithms and problem-solving training in C/C++ (Phitron)."),
        ("CSE Laboratory Training: ", "Completed university practical lab certifications in OOP, Database Management Systems, and Computer Networks.")
    ]
    for h_lbl, h_val in honors_trainings:
        p_h = doc.add_paragraph()
        p_h.paragraph_format.space_after = Pt(1.5)
        p_h.paragraph_format.line_spacing = 1.05
        r_hl = p_h.add_run(f"•  {h_lbl}")
        r_hl.font.bold = True
        r_hl.font.size = Pt(8.8)
        r_hl.font.color.rgb = RGBColor(15, 76, 129)
        r_hv = p_h.add_run(h_val)
        r_hv.font.size = Pt(8.6)

    # =========================================================================
    # 8. PERSONAL INFORMATION
    # =========================================================================
    add_section_header(doc, "8. Personal Information")
    
    personal_table = doc.add_table(rows=4, cols=4)
    personal_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    personal_table.autofit = False
    set_table_borders(personal_table, color="E2E8F0", sz="4")
    
    p_widths = [Inches(1.6), Inches(2.1), Inches(1.4), Inches(1.87)]
    for row in personal_table.rows:
        for idx, width in enumerate(p_widths):
            row.cells[idx].width = width
            
    p_info = [
        [("Father's Name", ": Md. Norul Amin"), ("Nationality", ": Bangladeshi (By Birth)")],
        [("Mother's Name", ": Nurun Naher"), ("Religion", ": Islam")],
        [("Date of Birth", ": 26/02/2004"), ("Marital Status", ": Single")],
        [("Permanent Address", ": Barmi, Sreepur, Gazipur-1743"), ("Present Address", ": Barmi, Sreepur, Gazipur-1743")]
    ]
    
    for row_idx, row_pair in enumerate(p_info):
        c0 = personal_table.rows[row_idx].cells[0]
        c1 = personal_table.rows[row_idx].cells[1]
        c2 = personal_table.rows[row_idx].cells[2]
        c3 = personal_table.rows[row_idx].cells[3]
        
        for c in [c0, c1, c2, c3]:
            set_cell_margins(c, top=30, bottom=30, left=40, right=40)
            
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(row_pair[0][0])
        r0.font.bold = True
        r0.font.size = Pt(8.5)
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(row_pair[0][1])
        r1.font.size = Pt(8.5)
        
        p2 = c2.paragraphs[0]
        p2.paragraph_format.space_after = Pt(0)
        r2 = p2.add_run(row_pair[1][0])
        r2.font.bold = True
        r2.font.size = Pt(8.5)
        
        p3 = c3.paragraphs[0]
        p3.paragraph_format.space_after = Pt(0)
        r3 = p3.add_run(row_pair[1][1])
        r3.font.size = Pt(8.5)

    # =========================================================================
    # 9. ACADEMIC & PROFESSIONAL REFERENCES
    # =========================================================================
    add_section_header(doc, "9. Academic & Professional References")
    
    ref_table = doc.add_table(rows=1, cols=2)
    ref_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    ref_table.autofit = False
    
    for idx, width in enumerate([Inches(3.48), Inches(3.49)]):
        ref_table.rows[0].cells[idx].width = width
        set_cell_margins(ref_table.rows[0].cells[idx], top=40, bottom=40, left=40, right=40)
        
    c_ref1 = ref_table.rows[0].cells[0]
    p_rf1 = c_ref1.paragraphs[0]
    p_rf1.paragraph_format.space_after = Pt(1)
    p_rf1.paragraph_format.line_spacing = 1.05
    r_rf1_n = p_rf1.add_run("Tanzillah Wahid\n")
    r_rf1_n.font.bold = True
    r_rf1_n.font.size = Pt(9)
    r_rf1_n.font.color.rgb = RGBColor(15, 76, 129)
    r_rf1_d = p_rf1.add_run(
        "Assistant Professor and Coordinator\n"
        "Department of Computer Science & Engineering\n"
        "Uttara University, Dhaka, Bangladesh\n"
        "Email: tanzillah@uttarauniversity.edu.bd"
    )
    r_rf1_d.font.size = Pt(8.4)
    
    c_ref2 = ref_table.rows[0].cells[1]
    p_rf2 = c_ref2.paragraphs[0]
    p_rf2.paragraph_format.space_after = Pt(1)
    p_rf2.paragraph_format.line_spacing = 1.05
    r_rf2_n = p_rf2.add_run("Md. Norul Amin\n")
    r_rf2_n.font.bold = True
    r_rf2_n.font.size = Pt(9)
    r_rf2_n.font.color.rgb = RGBColor(15, 76, 129)
    r_rf2_d = p_rf2.add_run(
        "Founder & Principal\n"
        "Barmi Al-Madina Pre-Cadet School\n"
        "Barmi, Sreepur, Gazipur, Bangladesh\n"
        "Phone: +880 1915-430867"
    )
    r_rf2_d.font.size = Pt(8.4)

    # =========================================================================
    # DECLARATION & SIGNATURE
    # =========================================================================
    p_dec = doc.add_paragraph()
    p_dec.paragraph_format.space_before = Pt(8)
    p_dec.paragraph_format.space_after = Pt(10)
    p_dec.paragraph_format.line_spacing = 1.1
    r_dec = p_dec.add_run(
        "Declaration: I, Nurul Arefin Nabil, solemnly declare that all the information provided in this curriculum vitae "
        "is authentic, complete, and accurate to the best of my knowledge and belief."
    )
    r_dec.font.italic = True
    r_dec.font.size = Pt(8.2)
    r_dec.font.color.rgb = RGBColor(71, 85, 105)
    
    # Signature line
    p_sig = doc.add_paragraph()
    p_sig.paragraph_format.space_after = Pt(0)
    p_sig.paragraph_format.keep_with_next = True
    r_sig_line = p_sig.add_run("_______________________________\n")
    r_sig_line.font.bold = True
    r_sig_line.font.color.rgb = RGBColor(15, 76, 129)
    r_sig_name = p_sig.add_run("Nurul Arefin Nabil\n")
    r_sig_name.font.bold = True
    r_sig_name.font.size = Pt(9.2)
    r_sig_date = p_sig.add_run("Date: .......................................")
    r_sig_date.font.size = Pt(8.3)
    r_sig_date.font.color.rgb = RGBColor(100, 116, 139)
    
    # =========================================================================
    # PAGE 3: TRADITIONAL / OLD-STYLE CV (STANDALONE 1-PAGE FORMAT)
    # =========================================================================
    add_old_style_cv_page(doc)
    
    output_filename = "CV_Nurul_Arefin_Nabil_ICT_Lecturer.docx"
    try:
        doc.save(output_filename)
        print(f"Successfully generated clean Word CV: {output_filename}")
    except PermissionError:
        alt_filename = "CV_Nurul_Arefin_Nabil_ICT_Teacher.docx"
        doc.save(alt_filename)
        print(f"Original file is open in Word. Saved to: {alt_filename}")

def add_old_style_cv_page(doc):
    # Add page break to start Page 3 (Traditional 1-page Old Style CV)
    doc.add_page_break()
    
    # Header Table: Title & Contacts on Left/Center, Photo Box on Right
    hdr_tbl = doc.add_table(rows=1, cols=2)
    hdr_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_tbl.autofit = False
    
    hdr_widths = [Inches(5.55), Inches(1.4)]
    for row in hdr_tbl.rows:
        for idx, w in enumerate(hdr_widths):
            row.cells[idx].width = w
            
    c_info = hdr_tbl.rows[0].cells[0]
    c_photo = hdr_tbl.rows[0].cells[1]
    set_cell_margins(c_info, top=0, bottom=15, left=0, right=30)
    set_cell_margins(c_photo, top=8, bottom=8, left=8, right=8)
    c_photo.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    
    # Border around photo cell for classic passport frame
    tcPr = c_photo._element.get_or_add_tcPr()
    p_bdr = parse_xml(f'''
    <w:tcBorders {nsdecls("w")}>
        <w:top w:val="single" w:sz="6" w:space="0" w:color="94A3B8"/>
        <w:left w:val="single" w:sz="6" w:space="0" w:color="94A3B8"/>
        <w:bottom w:val="single" w:sz="6" w:space="0" w:color="94A3B8"/>
        <w:right w:val="single" w:sz="6" w:space="0" w:color="94A3B8"/>
    </w:tcBorders>
    ''')
    tcPr.append(p_bdr)
    
    # Traditional Title
    p_t1 = c_info.paragraphs[0]
    p_t1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t1.paragraph_format.space_before = Pt(0)
    p_t1.paragraph_format.space_after = Pt(1)
    r_cv = p_t1.add_run("CURRICULUM VITAE\n")
    r_cv.font.name = 'Calibri'
    r_cv.font.size = Pt(13)
    r_cv.font.bold = True
    r_cv.font.underline = True
    r_cv.font.color.rgb = RGBColor(15, 76, 129)
    
    r_of = p_t1.add_run("OF\n")
    r_of.font.name = 'Calibri'
    r_of.font.size = Pt(8.5)
    r_of.font.bold = True
    
    r_name = p_t1.add_run("NURUL AREFIN NABIL")
    r_name.font.name = 'Calibri'
    r_name.font.size = Pt(13)
    r_name.font.bold = True
    r_name.font.color.rgb = RGBColor(15, 76, 129)
    
    # Contact info in traditional style
    p_cnt = c_info.add_paragraph()
    p_cnt.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cnt.paragraph_format.space_before = Pt(2)
    p_cnt.paragraph_format.space_after = Pt(0)
    p_cnt.paragraph_format.line_spacing = 1.05
    r_c = p_cnt.add_run(
        "Mailing Address: Barmi, Sreepur, Gazipur-1743, Dhaka, Bangladesh\n"
        "Mobile: +880 1881-196146   |   E-mail: nurularefinnabil@gmail.com\n"
        "Portfolio: www.areefin.me   |   LinkedIn: in/n-arefin-nabil"
    )
    r_c.font.size = Pt(8.2)
    r_c.font.color.rgb = RGBColor(51, 65, 85)
    
    # Photo insertion
    p_ph = c_photo.paragraphs[0]
    p_ph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ph.paragraph_format.space_after = Pt(0)
    img_path = "profile.jpg"
    if os.path.exists(img_path):
        p_ph.add_run().add_picture(img_path, width=Inches(1.15))
        
    # -------------------------------------------------------------------------
    # 1. CAREER OBJECTIVE
    # -------------------------------------------------------------------------
    add_classic_section_header(doc, "1. Career Objective")
    p_obj = doc.add_paragraph()
    p_obj.paragraph_format.space_before = Pt(2)
    p_obj.paragraph_format.space_after = Pt(3)
    p_obj.paragraph_format.line_spacing = 1.08
    p_obj.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r_obj = p_obj.add_run(
        "To secure a responsible position as an ICT / Computer Science Teacher or IT Professional in a reputed institution, "
        "utilizing my academic foundation in CSE, practical programming skills, and teaching capabilities to contribute "
        "effectively toward institutional excellence and students' academic growth."
    )
    r_obj.font.size = Pt(8.6)
    
    # -------------------------------------------------------------------------
    # 2. WORK & TEACHING EXPERIENCE
    # -------------------------------------------------------------------------
    add_classic_section_header(doc, "2. Work & Teaching Experience")
    
    exp_items = [
        ("ICT & Programming Instructor / Academic Tutor", "Tongi, Uttara & Sreepur", "2024 – Present (1+ Year)",
         "Teaching HSC ICT curriculum (C Programming, HTML/Web Design, Database/SQL, Number Systems). Conducting practical lab sessions and exam preparation model tests."),
        ("Assistant Teacher (ICT & Science)", "Barmi Al-Madina School & Plus Coaching", "1+ Year",
         "Conducted secondary ICT & Science classes, structured lesson planning, student evaluation, computer lab training, and typing coaching.")
    ]
    
    for title, loc, duration, desc in exp_items:
        p_e = doc.add_paragraph()
        p_e.paragraph_format.space_before = Pt(2)
        p_e.paragraph_format.space_after = Pt(1)
        p_e.paragraph_format.line_spacing = 1.05
        
        r_t = p_e.add_run(f"•  {title}")
        r_t.font.bold = True
        r_t.font.size = Pt(8.8)
        r_t.font.color.rgb = RGBColor(15, 76, 129)
        
        r_l = p_e.add_run(f" — {loc}")
        r_l.font.size = Pt(8.4)
        
        r_d = p_e.add_run(f"  ({duration})\n")
        r_d.font.bold = True
        r_d.font.size = Pt(8.2)
        r_d.font.color.rgb = RGBColor(100, 116, 139)
        
        r_ds = p_e.add_run(f"   {desc}")
        r_ds.font.size = Pt(8.3)
        r_ds.font.color.rgb = RGBColor(51, 65, 85)

    # -------------------------------------------------------------------------
    # 3. ACADEMIC QUALIFICATIONS (CLASSIC FULL GRID TABLE)
    # -------------------------------------------------------------------------
    add_classic_section_header(doc, "3. Educational Qualifications")
    
    edu_tbl = doc.add_table(rows=4, cols=5)
    edu_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    edu_tbl.autofit = False
    set_classic_grid_borders(edu_tbl, color="94A3B8", sz="4")
    
    edu_widths = [Inches(1.95), Inches(2.2), Inches(1.05), Inches(0.65), Inches(1.1)]
    for row in edu_tbl.rows:
        for idx, w in enumerate(edu_widths):
            row.cells[idx].width = w
            
    hdrs = ["Exam / Degree", "Board / University", "Group / Major", "Year", "Result"]
    for idx, text in enumerate(hdrs):
        cell = edu_tbl.rows[0].cells[idx]
        set_cell_background(cell, "F1F5F9")
        set_cell_margins(cell, top=45, bottom=45, left=50, right=50)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        if idx in [3, 4]:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.font.bold = True
        r.font.size = Pt(8.4)
        r.font.color.rgb = RGBColor(15, 76, 129)
        
    edu_rows = [
        ("B.Sc. in Computer Science & Engineering", "Uttara University, Dhaka", "CSE", "2026", "CGPA 3.40 / 4.00"),
        ("Alim (Equivalent to HSC)", "Tamirul Millat Kamil Madrasha", "Science", "2021", "GPA 5.00 / 5.00"),
        ("Dakhil (Equivalent to SSC)", "Patka Dakhil Madrasha", "Science", "2019", "GPA 5.00 (Upazila 1st)")
    ]
    for r_idx, data in enumerate(edu_rows, start=1):
        for c_idx, val in enumerate(data):
            c = edu_tbl.rows[r_idx].cells[c_idx]
            set_cell_margins(c, top=35, bottom=35, left=50, right=50)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if c_idx in [3, 4]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(val)
            r.font.size = Pt(8.2)
            if c_idx == 0:
                r.font.bold = True
            if "5.00" in val:
                r.font.bold = True

    # -------------------------------------------------------------------------
    # 4. COMPUTER & TECHNICAL SKILLS
    # -------------------------------------------------------------------------
    add_classic_section_header(doc, "4. Computer & Technical Skills")
    skills = [
        ("Office Applications", "MS Word, MS Excel, MS PowerPoint, MS Access, Internet & Google Workspace"),
        ("Programming & Web", "C, C++, Python, HTML5, CSS3, Relational Database & SQL queries"),
        ("Hardware & Systems", "Windows, Linux, Computer Troubleshooting, Lab Networking & Safe Internet")
    ]
    for cat, items in skills:
        p_s = doc.add_paragraph()
        p_s.paragraph_format.space_before = Pt(1)
        p_s.paragraph_format.space_after = Pt(1)
        p_s.paragraph_format.line_spacing = 1.05
        r_k = p_s.add_run(f"•  {cat}: ")
        r_k.font.bold = True
        r_k.font.size = Pt(8.4)
        r_v = p_s.add_run(items)
        r_v.font.size = Pt(8.3)

    # -------------------------------------------------------------------------
    # 5. PERSONAL DETAILS (TRADITIONAL 2-COLUMN TABLE FORMAT)
    # -------------------------------------------------------------------------
    add_classic_section_header(doc, "5. Personal Information")
    
    p_tbl = doc.add_table(rows=5, cols=4)
    p_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    p_tbl.autofit = False
    
    p_widths = [Inches(1.5), Inches(2.2), Inches(1.3), Inches(1.95)]
    for row in p_tbl.rows:
        for idx, w in enumerate(p_widths):
            row.cells[idx].width = w
            
    personal_data = [
        ("Father's Name", ": Md. Norul Amin", "Date of Birth", ": 26th February, 2004"),
        ("Mother's Name", ": Nurun Naher", "Gender / Sex", ": Male"),
        ("Nationality", ": Bangladeshi (by birth)", "Marital Status", ": Single"),
        ("Religion", ": Islam", "Blood Group", ": B+ (Positive)"),
        ("Permanent Address", ": Barmi, Sreepur, Gazipur", "Present Address", ": Barmi, Sreepur, Gazipur")
    ]
    
    for r_i, (k1, v1, k2, v2) in enumerate(personal_data):
        cells = p_tbl.rows[r_i].cells
        for c in cells:
            set_cell_margins(c, top=20, bottom=20, left=30, right=30)
            
        p0 = cells[0].paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(k1)
        r0.font.bold = True
        r0.font.size = Pt(8.2)
        
        p1 = cells[1].paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(v1)
        r1.font.size = Pt(8.2)
        
        p2 = cells[2].paragraphs[0]
        p2.paragraph_format.space_after = Pt(0)
        r2 = p2.add_run(k2)
        r2.font.bold = True
        r2.font.size = Pt(8.2)
        
        p3 = cells[3].paragraphs[0]
        p3.paragraph_format.space_after = Pt(0)
        r3 = p3.add_run(v2)
        r3.font.size = Pt(8.2)

    # -------------------------------------------------------------------------
    # 6. REFERENCES
    # -------------------------------------------------------------------------
    add_classic_section_header(doc, "6. References")
    
    r_tbl = doc.add_table(rows=1, cols=2)
    r_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    r_tbl.autofit = False
    for idx, w in enumerate([Inches(3.48), Inches(3.47)]):
        r_tbl.rows[0].cells[idx].width = w
        set_cell_margins(r_tbl.rows[0].cells[idx], top=20, bottom=20, left=30, right=30)
        
    c_r1 = r_tbl.rows[0].cells[0]
    p_r1 = c_r1.paragraphs[0]
    p_r1.paragraph_format.space_after = Pt(0)
    p_r1.paragraph_format.line_spacing = 1.05
    r_r1_n = p_r1.add_run("Tanzillah Wahid\n")
    r_r1_n.font.bold = True
    r_r1_n.font.size = Pt(8.5)
    r_r1_n.font.color.rgb = RGBColor(15, 76, 129)
    r_r1_d = p_r1.add_run("Assistant Professor & Coordinator, Dept. of CSE\nUttara University, Dhaka | Email: tanzillah@uttarauniversity.edu.bd")
    r_r1_d.font.size = Pt(8.0)
    
    c_r2 = r_tbl.rows[0].cells[1]
    p_r2 = c_r2.paragraphs[0]
    p_r2.paragraph_format.space_after = Pt(0)
    p_r2.paragraph_format.line_spacing = 1.05
    r_r2_n = p_r2.add_run("Md. Norul Amin\n")
    r_r2_n.font.bold = True
    r_r2_n.font.size = Pt(8.5)
    r_r2_n.font.color.rgb = RGBColor(15, 76, 129)
    r_r2_d = p_r2.add_run("Founder & Principal, Barmi Al-Madina Pre-Cadet School\nBarmi, Sreepur, Gazipur | Mobile: +880 1915-430867")
    r_r2_d.font.size = Pt(8.0)

    # -------------------------------------------------------------------------
    # DECLARATION & SIGNATURE
    # -------------------------------------------------------------------------
    p_dec = doc.add_paragraph()
    p_dec.paragraph_format.space_before = Pt(4)
    p_dec.paragraph_format.space_after = Pt(4)
    p_dec.paragraph_format.line_spacing = 1.05
    r_dec = p_dec.add_run(
        "Declaration: I do hereby declare that all information stated above is authentic, complete and true to the best of my knowledge."
    )
    r_dec.font.italic = True
    r_dec.font.size = Pt(8.0)
    r_dec.font.color.rgb = RGBColor(71, 85, 105)
    
    # Signature row
    sig_tbl = doc.add_table(rows=1, cols=2)
    sig_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_tbl.autofit = False
    for idx, w in enumerate([Inches(3.5), Inches(3.45)]):
        sig_tbl.rows[0].cells[idx].width = w
        set_cell_margins(sig_tbl.rows[0].cells[idx], top=8, bottom=0, left=0, right=0)
        
    p_dt = sig_tbl.rows[0].cells[0].paragraphs[0]
    p_dt.paragraph_format.space_after = Pt(0)
    r_dt = p_dt.add_run("\nDate: .......................................")
    r_dt.font.size = Pt(8.2)
    r_dt.font.color.rgb = RGBColor(100, 116, 139)
    
    p_sg = sig_tbl.rows[0].cells[1].paragraphs[0]
    p_sg.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sg.paragraph_format.space_after = Pt(0)
    r_sgl = p_sg.add_run("_______________________________\n")
    r_sgl.font.bold = True
    r_sgl.font.color.rgb = RGBColor(15, 76, 129)
    r_sgn = p_sg.add_run("(Nurul Arefin Nabil)")
    r_sgn.font.bold = True
    r_sgn.font.size = Pt(8.8)

if __name__ == "__main__":
    build_best_teacher_cv()
